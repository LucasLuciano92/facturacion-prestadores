"""Lectores, validación del lote y exportación de archivos de liquidación."""

import argparse
import csv
import json
from datetime import date
from decimal import Decimal
from pathlib import Path
import unicodedata

from openpyxl import load_workbook

from validaciones import normalizar_dni, normalizar_importe, validar_fecha


ALIAS = {
    "dni": "DNI", "documento": "DNI", "dni afiliado": "DNI",
    "fecha": "Fecha", "fecha atencion": "Fecha",
    "cod practica": "Cod_Practica", "codigo practica": "Cod_Practica",
    "importe": "Importe", "monto": "Importe",
}


def nombre_columna(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.lower().replace("_", " ").strip().split())


def leer_clinica(archivo):
    libro = load_workbook(archivo, read_only=True, data_only=True)
    try:
        filas = libro.active.iter_rows(values_only=True)
        encabezado = next(filas, None)
        if encabezado is None:
            raise ValueError("El Excel no tiene encabezado")
        columnas = [ALIAS.get(nombre_columna(c)) for c in encabezado]
        requeridas = {"DNI", "Fecha", "Cod_Practica", "Importe"}
        if not requeridas.issubset(set(columnas)):
            raise ValueError("Faltan columnas requeridas en el Excel")
        for nombre in requeridas:
            if columnas.count(nombre) != 1:
                raise ValueError("Columna repetida en el Excel: " + nombre)
        for numero, valores in enumerate(filas, start=2):
            if all(v is None for v in valores):
                continue
            fila = {c: v for c, v in zip(columnas, valores) if c}
            yield numero, fila, None
    finally:
        libro.close()


def leer_laboratorio(archivo):
    with Path(archivo).open(encoding="utf-8-sig") as entrada:
        for numero, linea in enumerate(entrada, start=1):
            linea = linea.rstrip("\r\n")
            if not linea.strip():
                continue
            if len(linea) != 34:
                yield numero, {"original": linea}, "Longitud TXT inválida"
                continue
            centavos = linea[24:34]
            if not centavos.isascii() or not centavos.isdigit():
                yield numero, {"original": linea}, "Importe TXT inválido"
                continue
            yield numero, {
                "DNI": linea[:8].strip(), "Fecha": linea[8:16],
                "Cod_Practica": linea[16:24].strip(),
                "Importe": Decimal(centavos) / 100,
            }, None


def leer_imagenes(archivo):
    with Path(archivo).open(encoding="utf-8-sig", newline="") as entrada:
        lector = csv.reader(entrada, delimiter=";")
        for numero, campos in enumerate(lector, start=1):
            if not campos or not any(c.strip() for c in campos):
                continue
            if numero == 1 and nombre_columna(campos[0]) in {"fecha", "fecha atencion"}:
                continue
            if len(campos) != 4:
                yield numero, {"original": campos}, "Cantidad de campos CSV inválida"
                continue
            yield numero, dict(zip(
                ["Fecha", "DNI", "Cod_Practica", "Importe"], campos
            )), None


def validar_registro(fila, prestador, formato, hoy, afiliados, vistos):
    try:
        dni = normalizar_dni(fila["DNI"])
        fecha = validar_fecha(fila["Fecha"], formato, hoy)
        importe = normalizar_importe(fila["Importe"])
        codigo = str(fila["Cod_Practica"] or "").strip()
        if not codigo:
            raise ValueError("Código de práctica vacío")
        if afiliados.get(dni) != "activo":
            raise ValueError("Afiliado inexistente o dado de baja")
        clave = (dni, fecha, codigo)
        if clave in vistos:
            raise ValueError("Práctica duplicada")
        registro = {"prestador": prestador, "dni": dni, "fecha": fecha,
                    "codigo": codigo, "importe": importe}
        vistos.add(clave)
        return registro, None
    except KeyError as error:
        return None, "Campo faltante: " + str(error.args[0])
    except ValueError as error:
        return None, str(error)


def procesar_clinica(fila, hoy, afiliados, vistos):
    return validar_registro(fila, "clinica", "%d/%m/%Y", hoy, afiliados, vistos)


def procesar_laboratorio(fila, hoy, afiliados, vistos):
    return validar_registro(fila, "laboratorio", "%Y%m%d", hoy, afiliados, vistos)


def procesar_imagenes(fila, hoy, afiliados, vistos):
    return validar_registro(fila, "imagenes", "%d-%m-%Y", hoy, afiliados, vistos)


LECTORES = {
    "clinica": (leer_clinica, procesar_clinica),
    "laboratorio": (leer_laboratorio, procesar_laboratorio),
    "imagenes": (leer_imagenes, procesar_imagenes),
}


def procesar_lote(archivos, afiliados, hoy):
    if not isinstance(afiliados, dict):
        raise ValueError("La base de afiliados no está disponible o es inválida")
    if any(estado not in {"activo", "baja"} for estado in afiliados.values()):
        raise ValueError("La base contiene estados de afiliación inválidos")
    validos, rechazos, vistos = [], [], set()
    for prestador, archivo in archivos:
        if prestador not in LECTORES:
            raise ValueError("Prestador sin lector implementado: " + prestador)
        lector, procesar = LECTORES[prestador]
        for numero, fila, error in lector(archivo):
            registro = None
            if error is None:
                registro, error = procesar(fila, hoy, afiliados, vistos)
            if error:
                rechazos.append({"prestador": prestador, "fila": numero,
                                 "motivo": error, "original": fila})
            else:
                validos.append(registro)
    return validos, rechazos


def exportar(validos, rechazos, carpeta):
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    with (carpeta / "liquidacion.csv").open("w", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(["prestador", "dni", "fecha", "codigo", "importe"])
        for r in validos:
            escritor.writerow([r["prestador"], r["dni"], r["fecha"].isoformat(),
                               r["codigo"], format(r["importe"], ".2f")])
    for prestador in LECTORES:
        with (carpeta / f"rechazos_{prestador}.csv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            escritor = csv.writer(f, delimiter=";")
            escritor.writerow(["prestador", "fila", "motivo", "original"])
            for r in rechazos:
                if r["prestador"] == prestador:
                    escritor.writerow([prestador, r["fila"], r["motivo"],
                                       json.dumps(r["original"], ensure_ascii=False, default=str)])
    resumen = {"recibidos": len(validos) + len(rechazos),
               "aceptados": len(validos), "rechazados": len(rechazos)}
    (carpeta / "resumen.json").write_text(
        json.dumps(resumen, indent=2) + "\n", encoding="utf-8"
    )
    return resumen


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--clinica", type=Path)
    parser.add_argument("--laboratorio", type=Path)
    parser.add_argument("--imagenes", type=Path)
    parser.add_argument("--afiliados", required=True, type=Path)
    parser.add_argument("--fecha", required=True, type=date.fromisoformat)
    parser.add_argument("--salida", type=Path, default=Path("salida"))
    args = parser.parse_args()
    archivos = [(p, getattr(args, p)) for p in LECTORES if getattr(args, p)]
    if not archivos:
        parser.error("Se requiere al menos un archivo de prestador")
    try:
        afiliados = json.loads(args.afiliados.read_text(encoding="utf-8"))
        validos, rechazos = procesar_lote(archivos, afiliados, args.fecha)
        print(json.dumps(exportar(validos, rechazos, args.salida), ensure_ascii=False))
    except (OSError, ValueError, csv.Error) as error:
        parser.exit(1, f"No se pudo procesar el lote: {error}\n")


if __name__ == "__main__":
    main()
